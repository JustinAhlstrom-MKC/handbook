#!/usr/bin/env python3
"""Generate the one-page handbook launch summary (Word .docx) with a QR code.

Output: output/MKC-Handbook-Launch-Summary.docx
The QR points to the on-site acknowledgement page so staff can read the
acknowledgement statements and sign.
"""
import os
import qrcode
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "output")
os.makedirs(OUT_DIR, exist_ok=True)

SITE_URL = "team.margies-kitchen.com"
ACK_URL = "https://team.margies-kitchen.com/10-acknowledgement/acknowledgement/"
LOGO = os.path.join(ROOT, "policies", "assets", "MKC_GreenBlock.png")
QR_PNG = os.path.join(OUT_DIR, "_ack_qr.png")
DOCX = os.path.join(OUT_DIR, "MKC-Handbook-Launch-Summary.docx")

GREEN = RGBColor(0x2E, 0x3A, 0x23)  # deep green to match brand block

# --- build QR ---
qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M,
                   box_size=10, border=2)
qr.add_data(ACK_URL)
qr.make(fit=True)
qr.make_image(fill_color="black", back_color="white").save(QR_PNG)

# --- helpers ---

def set_margins(section, inches):
    section.top_margin = section.bottom_margin = Inches(inches)
    section.left_margin = section.right_margin = Inches(inches)

def add_para(doc, text="", size=10.5, bold=False, italic=False, align=None,
             color=None, space_after=4, space_before=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return p

def add_bullet(doc, lead, rest, size=9.5):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(1.5)
    pf.space_before = Pt(0)
    r1 = p.add_run(lead)
    r1.bold = True
    r1.font.size = Pt(size)
    r1.font.color.rgb = GREEN
    r2 = p.add_run(rest)
    r2.font.size = Pt(size)
    return p

def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)

# --- document ---
doc = Document()
set_margins(doc.sections[0], 0.6)
# tighten default style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

# Logo
lp = doc.add_paragraph()
lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
lp.paragraph_format.space_after = Pt(6)
lp.add_run().add_picture(LOGO, width=Inches(2.3))

# Title + subtitle
add_para(doc, "Your Employee Handbook Is Here", size=24, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, color=GREEN, space_after=2)
add_para(doc, "Margie's Kitchen & Cocktails   •   Grackle", size=12,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# Intro
add_para(doc,
         "We've gathered all of our policies, guidelines, and team information into one "
         "easy-to-use online handbook. Whether you work at Margie's or Grackle, it's where "
         "to find how we do things, what we expect, and what you can count on from us — "
         "available anytime on your phone, tablet, or computer.",
         size=10.5, space_after=8)

# What's inside
add_para(doc, "What's inside", size=12.5, bold=True, color=GREEN, space_after=3)
sections = [
    ("Welcome & About Us", " — our story, mission, values, and who to contact"),
    ("Employment Basics", " — at-will employment, classifications, EEO, and accommodations"),
    ("Workplace Conduct", " — harassment, respect, confidentiality, and a safe environment"),
    ("Compensation & Time Off", " — pay, tips, breaks, PTO, sick & safe time, leave, and benefits"),
    ("Scheduling & Attendance", " — how schedules and time-off requests work"),
    ("Appearance & Professionalism", " — dress code and grooming standards"),
    ("Technology & Systems", " — Toast, 7shifts, and social media"),
    ("Safety & Security", " — food safety, fire safety, emergencies, and injury reporting"),
    ("Administrative", " — personnel files, discipline, and separation"),
    ("Acknowledgement", " — confirm you've received the handbook"),
]
for lead, rest in sections:
    add_bullet(doc, lead, rest)

# spacer
add_para(doc, "", space_after=4)

# Bottom: access + CTA (left) | QR (right)
table = doc.add_table(rows=1, cols=2)
no_borders(table)
table.columns[0].width = Inches(5.0)
table.columns[1].width = Inches(2.2)
left, right = table.rows[0].cells
left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

# left cell content
lc = left.paragraphs[0]
lc.paragraph_format.space_after = Pt(2)
r = lc.add_run("Read it & sign")
r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREEN

p = left.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Visit  ")
r.font.size = Pt(11)
r = p.add_run(SITE_URL)
r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREEN

p = left.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("or scan the code →  to review the handbook and submit your "
              "acknowledgement. ")
r.font.size = Pt(10.5)
r = p.add_run("Everyone needs to complete this.")
r.bold = True; r.font.size = Pt(10.5)

p = left.add_paragraph()
r = p.add_run("Questions? Talk to your manager, or reach out to Becky or Justin anytime.")
r.italic = True; r.font.size = Pt(9.5)

# right cell: QR + caption
qp = right.paragraphs[0]
qp.alignment = WD_ALIGN_PARAGRAPH.CENTER
qp.paragraph_format.space_after = Pt(0)
qp.add_run().add_picture(QR_PNG, width=Inches(1.7))
cap = right.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = cap.add_run("Scan to read & sign")
r.font.size = Pt(9); r.bold = True; r.font.color.rgb = GREEN

doc.save(DOCX)
print("Wrote", DOCX)
